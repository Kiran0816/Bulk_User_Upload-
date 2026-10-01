from datetime import datetime, date
from io import BytesIO
from flask import Flask, render_template, request, jsonify, send_file
from sqlalchemy.exc import IntegrityError
from openpyxl import load_workbook, Workbook
from database_setup import configure_database, User, db

def parse_datetime(val):
    if not val or val == '':
        return None
    if isinstance(val, (datetime, date)):
        return val
    if isinstance(val, str):
        val = val.strip()
        if not val:
            return None
        for fmt in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d', '%d/%m/%Y', '%Y/%m/%d %H:%M:%S', '%Y-%m-%dT%H:%M:%S'):
            try:
                return datetime.strptime(val, fmt)
            except ValueError:
                continue
    return None

def parse_int(val, default=None):
    if val is None or val == '':
        return default
    try:
        return int(float(val))
    except (ValueError, TypeError):
        return default

app = Flask(__name__)
configure_database(app)

# Route for the index
@app.route('/')
def index():
    return render_template('index.html')

# Route for uploading users
@app.route('/api/upload/users', methods=['POST'])
def upload_users():
    is_json_request = (
        request.headers.get('Accept') == 'application/json' or
        request.headers.get('X-Requested-With') == 'XMLHttpRequest' or
        request.args.get('format') == 'json'
    )

    def respond(data, status_code=200):
        if is_json_request:
            return jsonify(data), status_code
        return render_template('upload_result.html', response_data=data), status_code

    try:
        if 'file' not in request.files:
            return respond({'error': 'No file provided'}, 400)

        file = request.files['file']
        if not file or file.filename == '':
            return respond({'error': 'No file selected'}, 400)

        if not file.filename.lower().endswith(('.xlsx', '.xlsm', '.xltx', '.xltm')):
            return respond({'error': 'Invalid file format. Please upload an Excel (.xlsx) file.'}, 400)

        try:
            excel_data = load_workbook(file, data_only=True)
        except Exception as e:
            return respond({'error': f'Corrupted or unreadable Excel file: {str(e)}'}, 400)

        sheet = excel_data.active
        required_columns = [
            "username", "email", "passwordhash", "firstname", "lastname",
            "roleid", "departmentid", "isactive", "lastlogin", "createdat",
            "updatedat", "profilepictureurl", "phone", "timezone", "language"
        ]

        success_count = 0
        error_rows = []

        for row_index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            # Skip completely empty rows
            if not row or all(cell is None or str(cell).strip() == '' for cell in row):
                continue

            user_data = dict(zip(required_columns, row))

            if not user_data.get('email'):
                error_rows.append({'row_number': row_index, 'row_data': user_data, 'error': 'Missing required email'})
                continue

            if not user_data.get('username'):
                error_rows.append({'row_number': row_index, 'row_data': user_data, 'error': 'Missing required username'})
                continue

            existing_user = User.query.filter_by(email=user_data['email']).first()
            if existing_user:
                error_rows.append({'row_number': row_index, 'row_data': user_data, 'error': 'Duplicate email'})
                continue

            existing_username = User.query.filter_by(username=user_data['username']).first()
            if existing_username:
                error_rows.append({'row_number': row_index, 'row_data': user_data, 'error': 'Duplicate username'})
                continue

            try:
                user_data['lastlogin'] = parse_datetime(user_data.get('lastlogin'))
                user_data['createdat'] = parse_datetime(user_data.get('createdat'))
                user_data['updatedat'] = parse_datetime(user_data.get('updatedat'))
                user_data['roleid'] = parse_int(user_data.get('roleid'))
                user_data['departmentid'] = parse_int(user_data.get('departmentid'))
                user_data['isactive'] = parse_int(user_data.get('isactive'), default=1)

                new_user = User(**user_data)
                db.session.add(new_user)
                db.session.commit()
                success_count += 1

            except IntegrityError as e:
                db.session.rollback()
                error_rows.append({'row_number': row_index, 'row_data': user_data, 'error': str(e)})

        if error_rows:
            response_data = {
                'status': 'partial success with errors' if success_count > 0 else 'failed with errors',
                'success_count': success_count,
                'error_count': len(error_rows),
                'errors': error_rows
            }
        else:
            response_data = {
                'status': 'success',
                'success_count': success_count,
                'error_count': 0,
                'errors': []
            }

        status_code = 200 if (error_rows == [] or success_count > 0) else 400
        return respond(response_data, status_code)

    except FileNotFoundError as e:
        return respond({'error': 'File not found'}, 400)

    except Exception as e:
        app.logger.error(f'Error during bulk upload: {str(e)}')
        return respond({'error': f'Internal Server Error: {str(e)}'}, 500)

# ... (rest of your code)

# Route for downloading the template
@app.route('/api/download/template', methods=['GET'])
def download_template():
    try:
        wb = Workbook()
        ws = wb.active
        headers = ["username", "email", "passwordhash", "firstname", "lastname", "roleid", "departmentid", "isactive", "lastlogin", "createdat", "updatedat", "profilepictureurl", "phone", "timezone", "language"]
        ws.append(headers)
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        response = send_file(output, as_attachment=True, download_name="user_upload_template.xlsx")
        response.headers["Content-Disposition"] = "attachment; filename=user_upload_template.xlsx"
        return response

    except Exception as e:
        return jsonify({'error': str(e)}), 500

if __name__ == '__main__':
    with app.app_context():
        db.create_all()
    app.run(debug=True)
