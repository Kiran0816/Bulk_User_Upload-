import os
import io
import time
import pytest
from openpyxl import Workbook, load_workbook
from app import app, db, User

@pytest.fixture
def client():
    app.config['TESTING'] = True
    app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///:memory:'
    with app.app_context():
        db.create_all()
        yield app.test_client()
        db.session.remove()
        db.drop_all()

def create_sample_excel(records):
    wb = Workbook()
    ws = wb.active
    headers = [
        "username", "email", "passwordhash", "firstname", "lastname",
        "roleid", "departmentid", "isactive", "lastlogin", "createdat",
        "updatedat", "profilepictureurl", "phone", "timezone", "language"
    ]
    ws.append(headers)
    for rec in records:
        ws.append([rec.get(h, "") for h in headers])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

# =========================================================================
# 1. Template Download Tests
# =========================================================================

def test_download_template_status_and_headers(client):
    response = client.get('/api/download/template')
    assert response.status_code == 200
    assert "user_upload_template.xlsx" in response.headers.get("Content-Disposition", "")
    
    # Verify openpyxl can read the downloaded content
    wb = load_workbook(io.BytesIO(response.data))
    sheet = wb.active
    expected_headers = [
        "username", "email", "passwordhash", "firstname", "lastname",
        "roleid", "departmentid", "isactive", "lastlogin", "createdat",
        "updatedat", "profilepictureurl", "phone", "timezone", "language"
    ]
    actual_headers = [cell.value for cell in sheet[1]]
    assert actual_headers == expected_headers

# =========================================================================
# 2. Bulk Upload - Happy Path
# =========================================================================

def test_bulk_upload_valid_users(client):
    records = [
        {
            "username": f"user_{i}",
            "email": f"user_{i}@example.com",
            "passwordhash": "hash_xyz_123",
            "firstname": f"First{i}",
            "lastname": f"Last{i}",
            "roleid": 1,
            "departmentid": 2,
            "isactive": 1,
            "phone": "1234567890",
            "timezone": "UTC",
            "language": "en"
        }
        for i in range(1, 11)
    ]
    excel_file = create_sample_excel(records)
    response = client.post(
        '/api/upload/users',
        data={'file': (excel_file, 'test_users.xlsx')},
        content_type='multipart/form-data',
        headers={'Accept': 'application/json'}
    )
    assert response.status_code == 200
    data = response.get_json()
    assert data['status'] == 'success'
    assert data['success_count'] == 10
    assert data['error_count'] == 0

    with app.app_context():
        assert User.query.count() == 10

# =========================================================================
# 3. Duplicate Detection & Partial Success
# =========================================================================

def test_bulk_upload_duplicate_email_detection(client):
    # Seed 1 user first
    with app.app_context():
        u = User(username="existing", email="duplicate@example.com")
        db.session.add(u)
        db.session.commit()

    records = [
        {"username": "new_user_1", "email": "duplicate@example.com"},  # duplicate in DB
        {"username": "new_user_2", "email": "unique@example.com"}      # valid
    ]
    excel_file = create_sample_excel(records)
    response = client.post(
        '/api/upload/users',
        data={'file': (excel_file, 'test_duplicates.xlsx')},
        content_type='multipart/form-data',
        headers={'Accept': 'application/json'}
    )
    assert response.status_code == 200
    data = response.get_json()
    assert data['status'] == 'partial success with errors'
    assert data['success_count'] == 1
    assert data['error_count'] == 1
    assert data['errors'][0]['row_number'] == 2
    assert "Duplicate email" in data['errors'][0]['error']

def test_bulk_upload_duplicate_username_detection(client):
    with app.app_context():
        u = User(username="samename", email="first@example.com")
        db.session.add(u)
        db.session.commit()

    records = [
        {"username": "samename", "email": "second@example.com"}
    ]
    excel_file = create_sample_excel(records)
    response = client.post(
        '/api/upload/users',
        data={'file': (excel_file, 'test_user_dup.xlsx')},
        content_type='multipart/form-data',
        headers={'Accept': 'application/json'}
    )
    assert response.status_code == 400
    data = response.get_json()
    assert data['error_count'] == 1
    assert "Duplicate username" in data['errors'][0]['error']

# =========================================================================
# 4. Input Validation & Edge Cases
# =========================================================================

def test_bulk_upload_missing_email_and_username(client):
    records = [
        {"username": "", "email": "no_user@example.com"},
        {"username": "no_email", "email": ""}
    ]
    excel_file = create_sample_excel(records)
    response = client.post(
        '/api/upload/users',
        data={'file': (excel_file, 'missing_fields.xlsx')},
        content_type='multipart/form-data',
        headers={'Accept': 'application/json'}
    )
    assert response.status_code == 400
    data = response.get_json()
    assert data['error_count'] == 2

def test_bulk_upload_empty_rows_skipped(client):
    records = [
        {"username": "valid1", "email": "valid1@example.com"},
        {},  # empty row
        {"username": "valid2", "email": "valid2@example.com"},
        {}   # another empty row
    ]
    excel_file = create_sample_excel(records)
    response = client.post(
        '/api/upload/users',
        data={'file': (excel_file, 'empty_rows.xlsx')},
        content_type='multipart/form-data',
        headers={'Accept': 'application/json'}
    )
    assert response.status_code == 200
    data = response.get_json()
    assert data['success_count'] == 2
    assert data['error_count'] == 0

def test_no_file_provided(client):
    response = client.post(
        '/api/upload/users',
        data={},
        content_type='multipart/form-data',
        headers={'Accept': 'application/json'}
    )
    assert response.status_code == 400
    assert 'No file provided' in response.get_json()['error']

def test_empty_filename(client):
    response = client.post(
        '/api/upload/users',
        data={'file': (io.BytesIO(b''), '')},
        content_type='multipart/form-data',
        headers={'Accept': 'application/json'}
    )
    assert response.status_code == 400
    assert 'No file selected' in response.get_json()['error']

def test_invalid_file_extension(client):
    response = client.post(
        '/api/upload/users',
        data={'file': (io.BytesIO(b'col1,col2\nval1,val2'), 'users.csv')},
        content_type='multipart/form-data',
        headers={'Accept': 'application/json'}
    )
    assert response.status_code == 400
    assert 'Invalid file format' in response.get_json()['error']

def test_corrupted_excel_file(client):
    response = client.post(
        '/api/upload/users',
        data={'file': (io.BytesIO(b'Not an excel zip binary file!'), 'bad.xlsx')},
        content_type='multipart/form-data',
        headers={'Accept': 'application/json'}
    )
    assert response.status_code == 400
    assert 'Corrupted or unreadable' in response.get_json()['error']
