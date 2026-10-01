import io
import time
import statistics
from openpyxl import Workbook
from app import app, db, User

def generate_dataset(n_records, duplicate_ratio=0.0):
    wb = Workbook()
    ws = wb.active
    headers = [
        "username", "email", "passwordhash", "firstname", "lastname",
        "roleid", "departmentid", "isactive", "lastlogin", "createdat",
        "updatedat", "profilepictureurl", "phone", "timezone", "language"
    ]
    ws.append(headers)

    n_duplicates = int(n_records * duplicate_ratio)
    n_uniques = n_records - n_duplicates

    # Unique records
    for i in range(1, n_uniques + 1):
        ws.append([
            f"perf_user_{i}",
            f"perf_{i}@benchmark.com",
            f"argon2id$v=19$m=65536,t=3,p=4${i}",
            f"FirstName{i}",
            f"LastName{i}",
            1,
            2,
            1,
            "2026-01-01 12:00:00",
            "2026-01-01 12:00:00",
            "2026-01-01 12:00:00",
            f"https://cdn.example.com/avatar/{i}.png",
            f"+1555000{i:04d}",
            "America/New_York",
            "en-US"
        ])

    # Injected duplicate records
    for i in range(1, n_duplicates + 1):
        ws.append([
            f"dup_user_{i}",
            f"perf_1@benchmark.com", # Duplicate email intentional
            "pwd_hash_dup",
            "Duplicate",
            "Record",
            1, 2, 1, "", "", "", "", "", "UTC", "en"
        ])

    buf = io.BytesIO()
    wb.save(buf)
    file_size_bytes = buf.tell()
    buf.seek(0)
    return buf, file_size_bytes

def benchmark_batch_sizes(batch_sizes=[100, 500, 1000]):
    print("=" * 70)
    print("STARTING BULK USER UPLOAD EMPIRICAL BENCHMARK SUITE")
    print("=" * 70)

    results = []

    for size in batch_sizes:
        app.config['TESTING'] = True
        app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///:memory:'

        with app.app_context():
            db.create_all()
            client = app.test_client()

            excel_buf, file_size = generate_dataset(size, duplicate_ratio=0.05) # 5% intentional duplicates
            file_size_kb = file_size / 1024

            start_time = time.perf_counter()
            response = client.post(
                '/api/upload/users',
                data={'file': (excel_buf, f'bench_{size}.xlsx')},
                content_type='multipart/form-data',
                headers={'Accept': 'application/json'}
            )
            elapsed_time = time.perf_counter() - start_time

            data = response.get_json()
            success_count = data.get('success_count', 0)
            error_count = data.get('error_count', 0)
            throughput = size / elapsed_time
            ms_per_record = (elapsed_time / size) * 1000

            results.append({
                "size": size,
                "file_size_kb": file_size_kb,
                "elapsed_time_s": elapsed_time,
                "throughput_rps": throughput,
                "ms_per_record": ms_per_record,
                "success_count": success_count,
                "error_count": error_count,
                "status_code": response.status_code
            })

            print(f"Batch Size: {size:4d} rows | Payload: {file_size_kb:6.1f} KB | Time: {elapsed_time:6.3f}s | "
                  f"Throughput: {throughput:6.1f} rec/s | Avg Latency: {ms_per_record:5.2f} ms/rec | "
                  f"Success: {success_count} | Errs: {error_count}")

            db.session.remove()
            db.drop_all()

    print("=" * 70)
    return results

def benchmark_batch_commit_optimization(n_records=1000):
    print("\n" + "=" * 70)
    print(f"TRANSACTION OPTIMIZATION COMPARISON ({n_records} Records)")
    print("=" * 70)

    app.config['TESTING'] = True
    app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///:memory:'

    # 1. Unbatched (Row-by-Row commit - baseline)
    with app.app_context():
        db.create_all()
        excel_buf, _ = generate_dataset(n_records, duplicate_ratio=0.0)
        from openpyxl import load_workbook
        wb = load_workbook(excel_buf, data_only=True)
        sheet = wb.active
        headers = [
            "username", "email", "passwordhash", "firstname", "lastname",
            "roleid", "departmentid", "isactive", "lastlogin", "createdat",
            "updatedat", "profilepictureurl", "phone", "timezone", "language"
        ]

        t0 = time.perf_counter()
        for row in sheet.iter_rows(min_row=2, values_only=True):
            user_data = dict(zip(headers, row))
            user = User(username=user_data['username'], email=user_data['email'])
            db.session.add(user)
            db.session.commit() # Individual commit per row
        time_row_by_row = time.perf_counter() - t0
        throughput_row = n_records / time_row_by_row
        db.session.remove()
        db.drop_all()

    # 2. Batched (Single bulk transaction commit)
    with app.app_context():
        db.create_all()
        excel_buf, _ = generate_dataset(n_records, duplicate_ratio=0.0)
        wb = load_workbook(excel_buf, data_only=True)
        sheet = wb.active

        t0 = time.perf_counter()
        users_to_add = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            user_data = dict(zip(headers, row))
            users_to_add.append(User(username=user_data['username'], email=user_data['email']))
        db.session.bulk_save_objects(users_to_add)
        db.session.commit() # Single transaction commit
        time_batched = time.perf_counter() - t0
        throughput_batched = n_records / time_batched
        speedup = time_row_by_row / time_batched
        latency_reduction = ((time_row_by_row - time_batched) / time_row_by_row) * 100
        db.session.remove()
        db.drop_all()

    print(f"Row-by-Row (Baseline): {time_row_by_row:6.3f}s | Throughput: {throughput_row:6.1f} rec/s | Avg Latency: {(time_row_by_row/n_records)*1000:5.2f} ms")
    print(f"Batched Commit:        {time_batched:6.3f}s | Throughput: {throughput_batched:6.1f} rec/s | Avg Latency: {(time_batched/n_records)*1000:5.2f} ms")
    print(f"Performance Speedup:   {speedup:5.1f}x faster | Latency Reduction: {latency_reduction:.1f}%")
    print("=" * 70)

if __name__ == '__main__':
    benchmark_batch_sizes([50, 100, 250, 500, 1000])
    benchmark_batch_commit_optimization(1000)

